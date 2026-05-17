from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "交付"
OUT_PATH = OUT_DIR / "OCR识别接口图鉴信息站设计文档.docx"


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_para(doc, text, style=None, bold_lead=None):
    paragraph = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead) :])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.25
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    run = paragraph.add_run(text)
    set_run_font(run)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.2


def add_heading(doc, text, level):
    paragraph = doc.add_heading("", level=level)
    run = paragraph.add_run(text)
    set_run_font(run, size=18 if level == 1 else 14, bold=True, color=(23, 32, 42))
    return paragraph


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        run = hdr_cells[index].paragraphs[0].add_run(header)
        set_run_font(run, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            run = cells[index].paragraphs[0].add_run(value)
            set_run_font(run)
    doc.add_paragraph()
    return table


def build_doc():
    OUT_DIR.mkdir(exist_ok=True)
    doc = Document()

    styles = doc.styles
    styles["Normal"].font.name = "微软雅黑"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("OCR 识别接口图鉴信息站设计文档")
    set_run_font(title_run, size=22, bold=True, color=(23, 32, 42))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("面向接入方、业务人员与开发人员的 OCR 接口文档站")
    set_run_font(subtitle_run, size=12, color=(83, 98, 115))

    add_para(
        doc,
        "文档版本：v0.2  |  当前阶段：外部版页面与能力档案落地  |  资料口径：当前接口资料用于页面成品化展示，最终字段以后续正式接口文档校准。",
    )

    add_heading(doc, "一、项目定位", 1)
    add_para(
        doc,
        "本项目建设一个纯前端 OCR 接口信息站，用于向接入方集中展示识别能力、字段定义、返回结构、接口地址和阅读辅助说明。它不是在线调试平台，不提供上传入口，不承载真实样例，也不保存密钥或认证信息。",
    )
    add_para(
        doc,
        "站点的核心定位是“OCR 识别接口图鉴”。它保留接口文档的准确性和可维护性，同时用更清晰的信息架构、规范化示意样张和字段分组，让业务人员与开发人员都能快速理解接口能力。",
    )

    add_heading(doc, "二、设计理念与初衷", 1)
    add_para(
        doc,
        "1. 从“接口条目”升级为“识别能力档案”。传统 API 文档往往只把接口地址、请求参数、返回字段平铺出来，业务人员难以判断接口适用场景，开发人员也需要在密集表格中反复查找。OCR 接口天然对应证件、票据、凭证、文本版面等真实材料类型，因此更适合用“图鉴”和“档案”的方式组织。",
        bold_lead="1. 从“接口条目”升级为“识别能力档案”。",
    )
    add_para(
        doc,
        "2. 先让人理解，再让人对接。站点的第一目标不是炫技，而是降低理解成本。业务人员需要先知道“识别什么、可用于什么场景、能拿到哪些关键信息”；开发人员需要继续看到 endpoint、入参、返回结构、字段英文名、坐标和置信度等技术细节。双视角阅读就是为了解决同一份文档同时服务两类读者的问题。",
        bold_lead="2. 先让人理解，再让人对接。",
    )
    add_para(
        doc,
        "3. 用静态可视化替代重交互。由于站点不能提供真实样例和在线体验入口，视觉价值不应来自上传、调用或动画玩法，而应来自稳定、可审阅、可长期维护的静态可视化。示意样张和字段解剖图能帮助用户理解字段在材料中的位置，又不会暴露真实敏感信息。",
        bold_lead="3. 用静态可视化替代重交互。",
    )
    add_para(
        doc,
        "4. 可信边界优先。接口文档站的专业感首先来自可信：不展示真实身份证、真实发票、真实企业信息，不给出可误用的调试入口，不展示任何密钥或真实业务样例。当前资料可用于形成页面成品和能力档案，正式上线前再由最终接口文档统一校准字段口径。",
        bold_lead="4. 可信边界优先。",
    )
    add_para(
        doc,
        "5. 建立可复制的图文解释系统。真正能形成差异化的不是某个页面特效，而是一套能覆盖所有 OCR 类型的解释结构：接口概览、规范化示意样张、字段标注、字段分组、技术表格、错误码和相关接口。模板成熟后，扩展到更多接口才会稳定。",
        bold_lead="5. 建立可复制的图文解释系统。",
    )

    add_heading(doc, "三、资料与接口范围", 1)
    add_para(
        doc,
        "当前“其他”目录中的资料包括旧版 OCR 接口总文档、增值税发票服务、小票/登记证识别、铁路电子客票识别服务等 PDF。它们用于帮助设计覆盖实际存在的业务类型，但不作为最终发布版本的唯一依据。",
    )
    add_table(
        doc,
        ["资料", "用途", "设计处理"],
        [
            ["银商大脑-OCR-接口文档-0907.pdf", "提供 26 类接口的早期总览和字段结构", "用于能力索引、两个完整详情页和基础档案页的内容来源"],
            ["银联商务开放平台--增值税发票识别.pdf", "提供更近版本的增值税发票资料", "用于校准发票类接口措辞与字段覆盖方向"],
            ["银联商务开放平台--银商大脑OCR小票、登记证识别接口.pdf", "提供小票、机动车登记证、以旧换新回收凭单等参考", "补充能力索引，并形成对应基础档案页"],
            ["银联商务开放平台--铁路电子客票识别服务.pdf", "提供铁路电子客票识别服务参考", "补充能力索引，并形成对应基础档案页"],
        ],
    )

    add_heading(doc, "四、当前建设范围", 1)
    add_bullet(doc, "采用 Vite + React + TypeScript 单页应用，使用 react-router-dom 管理首页和详情页。")
    add_bullet(doc, "首页提供能力索引、分类筛选、关键词搜索、接口卡片矩阵。")
    add_bullet(doc, "详情页采用左侧目录、中间正文、右侧摘要卡片的文档阅读布局。")
    add_bullet(doc, "完整档案优先打磨两个接口：二代证（人像页+国徽页）和增值税发票。")
    add_bullet(doc, "其他接口形成基础能力档案页，展示接口用途、典型字段、统一请求参数、返回结构、错误码和相关接口。")
    add_bullet(doc, "后续拿到最终版接口文档后，以同一数据模型替换或补齐字段，无需重做页面结构。")

    add_heading(doc, "五、信息架构", 1)
    add_table(
        doc,
        ["页面/模块", "面向对象", "核心内容"],
        [
            ["首页顶部", "所有读者", "站点定位、覆盖材料类型、结构化返回和统一接口协议"],
            ["能力索引", "业务人员", "按证照、票据、车辆、商户、通用文本等分类浏览接口"],
            ["接口卡片", "业务人员 + 开发人员", "接口名称、分类、摘要、典型字段、档案完整度"],
            ["详情概览", "业务人员", "用途、场景、字段数量、请求方式、返回形式"],
            ["示意样张", "业务人员", "规范化材料外观和字段位置，不使用真实样例"],
            ["字段分组", "业务人员 + 开发人员", "按语义组织字段，同时保留英文字段名和类型"],
            ["技术区", "开发人员", "测试/生产地址、请求参数、通用返回结构、错误码"],
            ["相关接口", "所有读者", "同类能力横向跳转"],
        ],
    )

    add_heading(doc, "六、视觉设计原则", 1)
    add_bullet(doc, "整体风格为科技文档风：浅色背景、深色正文、蓝灰基础、冷青色点缀。")
    add_bullet(doc, "控制装饰：不用大面积营销式渐变，不做夸张动效，不堆叠卡片。")
    add_bullet(doc, "强调信息密度：表格紧凑但可读，字段标签清晰，长字段支持换行。")
    add_bullet(doc, "示意图采用规范化虚构内容：看得出材料类型，但明显不是可被误用的真实样张。")
    add_bullet(doc, "响应式优先：移动端改为单列阅读，表格横向滚动，避免文字遮挡和重叠。")

    add_heading(doc, "七、数据模型与维护方式", 1)
    add_para(
        doc,
        "站点采用数据驱动。接口内容集中维护在 TypeScript 数据文件中，UI 组件只负责展示。后续拿到最终接口文档后，只需要补充或替换数据对象，就能批量扩展详情页。",
    )
    add_table(
        doc,
        ["模型", "关键字段", "说明"],
        [
            ["OcrInterfaceDoc", "slug、title、category、summary、endpoints、request、fieldGroups、response、errors、visual", "单个完整接口详情页的数据来源"],
            ["FieldDef", "key、label、type、required、description、tags、formatHint、notes", "字段表、业务卡片和标签体系的统一字段定义"],
            ["EndpointInfo", "protocol、method、testUrl、prodUrl", "开发视角展示接口地址和方法"],
            ["VisualSpec", "kind、coverAlt、detailAlt、annotations", "驱动封面图、详情示意图和字段标注"],
        ],
    )

    add_heading(doc, "八、安全与内容边界", 1)
    add_bullet(doc, "不提供上传、调用、调试、复制密钥等在线能力。")
    add_bullet(doc, "不展示真实身份证、真实发票、真实企业、真实交易、真实个人敏感信息。")
    add_bullet(doc, "示例值只使用格式占位或虚构内容，并在页面中保留“示意”语义。")
    add_bullet(doc, "最终上线前以正式命名文件夹中的最新接口文档为准，对当前页面中的字段名、endpoint 和说明做一次口径校准。")

    add_heading(doc, "九、验收标准", 1)
    add_bullet(doc, "业务人员能在 30 秒内判断接口识别对象、适用场景和关键字段。")
    add_bullet(doc, "开发人员能快速找到 endpoint、请求参数、返回结构、字段英文名、坐标和置信度说明。")
    add_bullet(doc, "页面在桌面、平板和手机下不出现文字溢出、遮挡或元素重叠。")
    add_bullet(doc, "所有详情页都包含概览、示意样张、字段分组、请求参数、返回结构、错误码和相关接口。")
    add_bullet(doc, "完整档案与基础档案在视觉状态上有区分，但基础档案也应是可阅读、可交付的成品页面。")

    add_heading(doc, "十、后续扩展流程", 1)
    add_para(
        doc,
        "当最终版接口文档放入正式命名文件夹后，建议按以下流程扩展：先抽取接口清单与字段表，确认重复接口与版本优先级；再将字段映射为 OcrInterfaceDoc 数据对象；然后按接口类型生成规范化示意样张；最后进行内容校验、响应式检查和业务/开发双视角验收。",
    )

    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    print(build_doc())
